"""Exportação do histórico de movimentos do MotorATP para .xlsx.

Gera uma planilha com duas abas:

- "Movimentos": o histórico completo (uma linha por evento), com físico,
  disponível, reservas por canal e ATP por canal logo após cada movimento.
- "Configuração": o estoque físico inicial e a política (proteção/restrição)
  de cada canal.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from estoque_atp.motor import MotorATP

# Cores por tipo de evento (fundo da célula "Evento").
_CORES_EVENTO = {
    "INICIAL": "D9D9D9",       # cinza
    "RESERVA": "FFE699",       # amarelo
    "EFETIVACAO": "C6EFCE",    # verde
    "CANCELAMENTO": "FFC7CE",  # vermelho
}
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def exportar_xlsx(motor: MotorATP, caminho: str | Path) -> Path:
    """Escreve o histórico de `motor` no arquivo .xlsx em `caminho`.

    Retorna o caminho do arquivo gravado.
    """
    caminho = Path(caminho)
    wb = Workbook()

    _aba_movimentos(wb.active, motor)
    _aba_configuracao(wb.create_sheet("Configuração"), motor)

    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho


def _aba_movimentos(ws, motor: MotorATP) -> None:
    ws.title = "Movimentos"
    canais = list(motor.canais)

    cabecalho = (
        ["Seq", "Momento", "Evento", "Canal", "Qtde", "Rótulo", "Físico", "Disponível"]
        + [f"Reserva {c}" for c in canais]
        + [f"ATP {c}" for c in canais]
    )
    ws.append(cabecalho)
    for col in range(1, len(cabecalho) + 1):
        celula = ws.cell(row=1, column=col)
        celula.fill = _HEADER_FILL
        celula.font = _HEADER_FONT
        celula.alignment = Alignment(horizontal="center", vertical="center")

    for mov in motor.historico:
        linha = (
            [
                mov.seq,
                mov.momento.strftime("%Y-%m-%d %H:%M:%S"),
                mov.evento,
                mov.canal,
                mov.quantidade,
                mov.rotulo,
                mov.fisico,
                mov.disponivel,
            ]
            + [mov.reservas[c] for c in canais]
            + [mov.atp[c] for c in canais]
        )
        ws.append(linha)
        # Colore a célula "Evento" (coluna 3) conforme o tipo.
        cor = _CORES_EVENTO.get(mov.evento)
        if cor:
            ws.cell(row=ws.max_row, column=3).fill = PatternFill("solid", fgColor=cor)

    ws.freeze_panes = "A2"
    _ajustar_larguras(ws, cabecalho)


def _aba_configuracao(ws, motor: MotorATP) -> None:
    ws.append(["Parâmetro", "Valor"])
    ws.append(["Estoque físico inicial", motor.historico[0].fisico])
    ws.append([])
    ws.append(["Canal", "Proteção", "Restrição"])

    for canal in motor.canais.values():
        ws.append(
            [
                canal.nome,
                canal.protecao,
                "sem teto" if canal.restricao is None else canal.restricao,
            ]
        )

    for linha in (1, 4):
        for col in range(1, 4):
            celula = ws.cell(row=linha, column=col)
            if celula.value is not None:
                celula.fill = _HEADER_FILL
                celula.font = _HEADER_FONT
    _ajustar_larguras(ws, ["Canal", "Proteção", "Restrição"], minimo=14)


def _ajustar_larguras(ws, cabecalho: list[str], minimo: int = 8) -> None:
    for col in range(1, len(cabecalho) + 1):
        largura = max(minimo, len(str(cabecalho[col - 1])) + 2)
        for row in range(2, ws.max_row + 1):
            valor = ws.cell(row=row, column=col).value
            if valor is not None:
                largura = max(largura, len(str(valor)) + 2)
        ws.column_dimensions[get_column_letter(col)].width = largura
