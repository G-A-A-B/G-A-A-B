"""Testes do histórico de movimentos e da exportação para .xlsx."""

from __future__ import annotations

from openpyxl import load_workbook

from estoque_atp import Canal, MotorATP, exportar_xlsx


def motor_padrao() -> MotorATP:
    return MotorATP(
        fisico=100,
        canais=[
            Canal("Loja", protecao=30),
            Canal("Site"),
            Canal("Marketplace", restricao=20),
        ],
    )


# --------------------------------------------------------------------------- #
# Histórico em memória
# --------------------------------------------------------------------------- #
def test_estado_inicial_e_registrado():
    m = motor_padrao()
    assert len(m.historico) == 1
    inicial = m.historico[0]
    assert inicial.evento == "INICIAL"
    assert inicial.fisico == 100
    assert inicial.atp == {"Loja": 100, "Site": 70, "Marketplace": 20}


def test_cada_operacao_gera_um_movimento():
    m = motor_padrao()
    m.reservar("Marketplace", 15)
    m.reservar("Site", 40)
    m.efetivar("Marketplace", 15)
    m.cancelar("Site", 10)
    # 1 inicial + 4 operações
    assert [mv.evento for mv in m.historico] == [
        "INICIAL",
        "RESERVA",
        "RESERVA",
        "EFETIVACAO",
        "CANCELAMENTO",
    ]
    assert [mv.seq for mv in m.historico] == [0, 1, 2, 3, 4]


def test_movimento_captura_estado_apos_o_evento():
    m = motor_padrao()
    m.reservar("Site", 40, "reserva site")
    mov = m.historico[-1]
    assert mov.canal == "Site"
    assert mov.quantidade == 40
    assert mov.rotulo == "reserva site"
    assert mov.disponivel == 60
    assert mov.reservas["Site"] == 40


def test_operacao_recusada_nao_registra_movimento():
    m = motor_padrao()
    try:
        m.reservar("Marketplace", 999)  # excede o ATP
    except Exception:
        pass
    # Apenas o INICIAL permanece no histórico.
    assert len(m.historico) == 1


# --------------------------------------------------------------------------- #
# Exportação para .xlsx
# --------------------------------------------------------------------------- #
def test_exportar_xlsx_gera_arquivo_com_historico(tmp_path):
    m = motor_padrao()
    m.reservar("Marketplace", 15)
    m.reservar("Site", 40)
    m.efetivar("Marketplace", 15)

    destino = exportar_xlsx(m, tmp_path / "hist.xlsx")
    assert destino.exists()

    wb = load_workbook(destino)
    assert wb.sheetnames == ["Movimentos", "Configuração"]

    ws = wb["Movimentos"]
    # 1 cabeçalho + 4 movimentos (inicial + 3 operações)
    assert ws.max_row == 5

    cabecalho = [c.value for c in ws[1]]
    assert cabecalho[:6] == ["Seq", "Momento", "Evento", "Canal", "Qtde", "Rótulo"]
    assert "ATP Loja" in cabecalho
    assert "Reserva Marketplace" in cabecalho

    # Última linha reflete o estado final (após efetivar 15 do Marketplace).
    ultima = {cab: cel.value for cab, cel in zip(cabecalho, ws[ws.max_row])}
    assert ultima["Evento"] == "EFETIVACAO"
    assert ultima["Físico"] == 85
    assert ultima["ATP Marketplace"] == 15


def test_aba_configuracao_reflete_politicas(tmp_path):
    m = motor_padrao()
    destino = exportar_xlsx(m, tmp_path / "hist.xlsx")
    wb = load_workbook(destino)
    ws = wb["Configuração"]
    # Normaliza cada linha removendo células vazias à direita.
    linhas = []
    for row in ws.iter_rows(values_only=True):
        valores = list(row)
        while valores and valores[-1] is None:
            valores.pop()
        linhas.append(tuple(valores))
    assert ("Estoque físico inicial", 100) in linhas
    assert ("Loja", 30, "sem teto") in linhas
    assert ("Marketplace", 0, 20) in linhas
